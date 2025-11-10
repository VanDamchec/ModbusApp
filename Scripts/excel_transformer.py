# -*- coding: utf-8 -*-
"""
excel_transformer.py — универсальный Excel→Excel преобразователь (openpyxl)
---------------------------------------------------------------------------
Дополнения:
- Ссылки на столбцы: имена заголовков, буквы ("C","AA") или индексы (1=A).
- Работа без заголовков (header_row: null/0).
- Фильтры и шаблоны понимают {A}, {AA}, {Header}.
- NEW: вычисления по выражениям:
  * В шаблонах: {python <expr>}
  * В маппингах:  {"to": "...", "expr": "<expr>"}
  В выражениях доступны:
    - переменные по буквам столбцов (A, B, C, ...),
    - переменные по заголовкам (нормализованные в валидные идентификаторы),
    - словарь V с исходными ключами (V["Адрес"], V["Address"], V["A"] ...),
    - модуль re, функции: str,int,float,len,abs,round,min,max.
  Пример:
    {"to": "Name", "expr": " (K or '').split(':')[0] if 'BITMAP' in (L or '') else K "}
    {"to": "Name", "template": "{J}_{python (K or '').split(':')[0] if 'BITMAP' in (L or '') else K}"}
"""
import argparse
import json
import re
from typing import Any, Dict, List, Optional, Union
from openpyxl import load_workbook, Workbook

Value = Union[str, int, float, None]
ColRef = Union[str, int]

# ----------------- utils -----------------
_LETTERS_RE = re.compile(r'^[A-Za-z]{1,3}$')  # ограничим до A..ZZZ, чтобы ловить опечатки
IDENT_RE = re.compile(r'\W+')  # всё, что не [A-Za-z0-9_]

def col_letter_to_index(s: str) -> Optional[int]:
    s = s.strip().upper()
    if not _LETTERS_RE.match(s):
        return None
    n = 0
    for ch in s:
        n = n * 26 + (ord(ch) - ord('A') + 1)
    return n

def parse_col_ref(ref: ColRef, headers_map: Optional[Dict[str, int]] = None) -> Optional[int]:
    if isinstance(ref, int):
        return ref
    if not isinstance(ref, str):
        return None
    if headers_map is not None and ref in headers_map:
        return headers_map[ref]
    idx = col_letter_to_index(ref)
    if idx is not None:
        return idx
    try:
        return int(ref)
    except Exception:
        return None

def index_to_col_letter(idx: int) -> str:
    res = []
    x = idx
    while x > 0:
        x, r = divmod(x - 1, 26)
        res.append(chr(ord('A') + r))
    return ''.join(reversed(res))

def _normalize_header_map(ws, header_row: Optional[int]) -> Dict[str, int]:
    if not header_row or header_row <= 0:
        return {}
    headers: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=col).value
        if v is None:
            continue
        name = str(v).strip()
        if name:
            headers[name] = col
    return headers

def _get_cell(ws, row: int, col_idx: int) -> Value:
    return ws.cell(row=row, column=col_idx).value

def _as_nonempty(v: Any) -> bool:
    return v is not None and str(v).strip() != ''

def _sanitize_ident(name: str) -> str:
    name = str(name)
    ident = IDENT_RE.sub('_', name).strip('_')
    if not ident:
        ident = '_'
    if ident[0].isdigit():
        ident = '_' + ident
    return ident

def _collect_row_lookups(ws, row: int, headers_map: Dict[str, int]) -> Dict[str, Any]:
    vals: Dict[str, Any] = {}
    for name, col in headers_map.items():
        vals[name] = _get_cell(ws, row, col)
    for col in range(1, ws.max_column + 1):
        letter = index_to_col_letter(col)
        vals[letter] = _get_cell(ws, row, col)
    return vals

# ----------------- шаблоны и выражения -----------------
_PY_PLACEHOLDER_RE = re.compile(r'\{python\s+([^}]+)\}')

def _build_env(row_lookup: Dict[str, Any]) -> Dict[str, Any]:
    env: Dict[str, Any] = {}
    for k, v in row_lookup.items():
        env[_sanitize_ident(k)] = v
        if isinstance(k, str) and _LETTERS_RE.match(k):
            env[k] = v
    env['V'] = row_lookup
    env.update({
        're': re,
        'str': str,
        'int': int,
        'float': float,
        'len': len,
        'abs': abs,
        'round': round,
        'min': min,
        'max': max,
    })
    return env

def _safe_eval(expr: str, env: Dict[str, Any]) -> Any:
    code = compile(expr, '<expr>', 'eval')
    return eval(code, {'__builtins__': {}}, env)

def _format_template(template: str, row_lookup: Dict[str, Any], env: Dict[str, Any]) -> str:
    out = template
    for k, v in row_lookup.items():
        out = out.replace('{%s}' % k, '' if v is None else str(v))
    def _repl(m):
        expr = m.group(1)
        try:
            val = _safe_eval(expr, env)
            return '' if val is None else str(val)
        except Exception:
            return ''
    out = _PY_PLACEHOLDER_RE.sub(_repl, out)
    out = re.sub(r'\{[^}]+\}', '', out)
    return out

# ----------------- фильтры -----------------
def _row_passes_filter(row_lookup: Dict[str, Any], require_all, require_any) -> bool:
    def lookup(colref):
        return row_lookup.get(str(colref))
    ok_all = True
    if require_all:
        ok_all = all(_as_nonempty(lookup(k)) for k in require_all)
    ok_any = True
    if require_any:
        ok_any = any(_as_nonempty(lookup(k)) for k in require_any)
    return ok_all and ok_any

# ----------------- основная трансформация -----------------
def transform_excel(input_path: str, output_path: str, config: Dict[str, Any]) -> bool:
    wb_src = load_workbook(input_path, data_only=True)
    ws_src = wb_src[config.get('source', {}).get('sheet')] if config.get('source', {}).get('sheet') else wb_src.active
    src_header_row_raw = config.get('source', {}).get('header_row', 1)
    src_header_row = src_header_row_raw if (isinstance(src_header_row_raw, int) and src_header_row_raw > 0) else None
    src_headers = _normalize_header_map(ws_src, src_header_row)

    wb_dst = Workbook()
    ws_dst = wb_dst.active
    ws_dst.title = config.get('target', {}).get('sheet', 'Sheet1')
    dst_header_row_raw = config.get('target', {}).get('header_row', 1)
    dst_header_row = dst_header_row_raw if (isinstance(dst_header_row_raw, int) and dst_header_row_raw > 0) else None

    target_headers: List[str] = config.get('target', {}).get('headers') or []
    write_headers = bool(target_headers) and bool(dst_header_row)

    if write_headers:
        for col_idx, name in enumerate(target_headers, start=1):
            ws_dst.cell(row=dst_header_row, column=col_idx, value=name)
        dst_index_by_name = {name: i+1 for i, name in enumerate(target_headers)}
    else:
        dst_index_by_name = {}

    filters_cfg = config.get('row_filter', {})
    require_all = filters_cfg.get('require_all', []) or []
    require_any = filters_cfg.get('require_any', []) or []

    start_src_row = (src_header_row + 1) if src_header_row else 1
    out_row = (dst_header_row + 1) if dst_header_row else 1
    last_row = ws_src.max_row

    for r in range(start_src_row, last_row + 1):
        if all(ws_src.cell(row=r, column=c).value in (None, '') for c in range(1, ws_src.max_column + 1)):
            continue

        row_lookup = _collect_row_lookups(ws_src, r, src_headers)
        if (require_all or require_any) and not _row_passes_filter(row_lookup, require_all, require_any):
            continue

        env = _build_env(row_lookup)

        for m in config.get('mappings', []):
            to_ref = m['to']
            dst_col = None
            if isinstance(to_ref, str) and to_ref in dst_index_by_name:
                dst_col = dst_index_by_name[to_ref]
            if dst_col is None:
                dst_col = parse_col_ref(to_ref, headers_map=None)
            if not dst_col or dst_col > 16384:
                raise ValueError(f'Unknown/invalid target column "{to_ref}". Use target.headers name, letter (A..ZZZ) or index (1..16384).')

            val: Any = None
            if 'value' in m:
                val = m['value']
            elif 'expr' in m:
                val = _safe_eval(m['expr'], env)
            elif 'template' in m:
                val = _format_template(m['template'], row_lookup, env)
            elif 'from' in m:
                src_ref = m['from']
                src_col = parse_col_ref(src_ref, headers_map=src_headers if src_headers else None)
                if src_col:
                    val = _get_cell(ws_src, r, src_col)
                else:
                    if isinstance(src_ref, str):
                        val = row_lookup.get(src_ref)
                    else:
                        val = None
            else:
                continue

            if m.get('skip_if_blank') and not _as_nonempty(val):
                continue

            if 'map' in m and isinstance(m['map'], dict):
                key = val
                if m.get('use_lowercase', False) and isinstance(key, str):
                    key = key.lower().strip()
                val = m['map'].get(key, m.get('default', val))

            ws_dst.cell(row=out_row, column=dst_col, value=val)

        out_row += 1

    wb_dst.save(output_path)
    return True

def main():
    ap = argparse.ArgumentParser(description='Excel→Excel универсальный преобразователь по JSON-конфигу (openpyxl).')
    ap.add_argument('--input', required=True, help='Путь к исходному .xlsx')
    ap.add_argument('--output', required=True, help='Путь для сохранения результата .xlsx (файл, не папка)')
    ap.add_argument('--config', required=False, help='Путь к config.json')
    ap.add_argument('--print-config-skeleton', action='store_true', help='Вывести черновик конфига по заголовкам источника')
    args = ap.parse_args()

    if args.print_config_skeleton:
        wb = load_workbook(args.input, data_only=True)
        ws = wb.active
        headers = {}
        header_row = 1
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if v is None:
                continue
            headers[str(v)] = c
        skeleton = {
            "source": {"sheet": ws.title, "header_row": 1},
            "target": {"sheet": "Sheet1", "header_row": 1, "headers": list(headers.keys())},
            "row_filter": {"require_all": [], "require_any": []},
            "mappings": [{"to": name, "from": name} for name in headers.keys()]
        }
        print(json.dumps(skeleton, ensure_ascii=False, indent=2))
        return

    if not args.config:
        raise SystemExit('Нужен --config путь к JSON-файлу. Или используйте --print-config-skeleton.')

    with open(args.config, 'r', encoding='utf-8') as f:
        cfg = json.load(f)

    transform_excel(args.input, args.output, cfg)

if __name__ == '__main__':
    main()
